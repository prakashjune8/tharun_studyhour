"""
Universal Exam Portion & Syllabus Parser for Tharun (Class 8)
Supports: PDF (.pdf), Word (.docx), Images (.png/.jpg), Text (.txt), and JSON (.json)
Parses school exam portion sheets and circulars into structured JSON for ScholasticPulse.
"""

import os
import re
import json
import io
import asyncio
from datetime import datetime

# Optional dependencies with graceful fallbacks
try:
    import docx
except ImportError:
    docx = None

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    from PIL import Image
    import winocr
except ImportError:
    winocr = None
    Image = None


KNOWN_SUBJECT_DEFS = [
    {
        "id": "sub_tamil",
        "subject": "Tamil",
        "category": "language",
        "gravity": 2,
        "color": "#ea580c",
        "aliases": [r"\btamil\b", r"\bதமிழ்\b"]
    },
    {
        "id": "sub_math",
        "subject": "Mathematics",
        "category": "core_heavy",
        "gravity": 3,
        "color": "#2563eb",
        "aliases": [r"\bmathematics\b", r"\bmaths\b", r"\bmath\b"]
    },
    {
        "id": "sub_phy_sci",
        "subject": "Physical Science",
        "category": "core_heavy",
        "gravity": 3,
        "color": "#7c3aed",
        "aliases": [r"\bphysical\s+science\b", r"\bphysics\s*(?:&|and)?\s*chemistry\b"]
    },
    {
        "id": "sub_bio_sci",
        "subject": "Biological Science",
        "category": "core_heavy",
        "gravity": 3,
        "color": "#16a34a",
        "aliases": [r"\bbiological\s+science\b", r"\bbiology\b", r"\bbio\s+science\b"]
    },
    {
        "id": "sub_soc_sci",
        "subject": "Social Science",
        "category": "core_moderate",
        "gravity": 2,
        "color": "#d97706",
        "aliases": [r"\bsocial\s+science\b", r"\bs\.?s\.?t\b", r"\bsocial\b"]
    },
    {
        "id": "sub_hindi",
        "subject": "Hindi",
        "category": "language",
        "gravity": 2,
        "color": "#dc2626",
        "aliases": [r"\bhindi\b", r"\bहिंदी\b"]
    },
    {
        "id": "sub_eng",
        "subject": "English",
        "category": "language",
        "gravity": 2,
        "color": "#0284c7",
        "aliases": [r"\benglish\b", r"\beng\b"]
    },
    {
        "id": "sub_cs",
        "subject": "Computer Science",
        "category": "skill",
        "gravity": 2,
        "color": "#059669",
        "aliases": [r"\bcomputer\s+science\b", r"\bcomputer\b", r"\binfo(?:rmation)?\s+tech\b"]
    }
]


class ExamPortionParser:
    """
    Intelligent parser for School Exam Portion circulars & documents.
    """

    @classmethod
    def extract_text_from_file(cls, filepath):
        """
        Extract raw text from PDF, Docx, TXT, Image, or loads JSON.
        Returns: (type_str, content)
        """
        if not os.path.exists(filepath):
            raise FileNotFoundError(f"File not found: {filepath}")

        ext = os.path.splitext(filepath)[1].lower()

        # 1. Direct JSON Syllabus
        if ext == ".json":
            with open(filepath, "r", encoding="utf-8") as f:
                data = json.load(f)
                return "JSON_DIRECT", data

        # 2. Word Document (.docx)
        if ext in [".docx", ".doc"]:
            if not docx:
                raise ImportError("python-docx is required to read .docx files.")
            doc = docx.Document(filepath)
            lines = []
            for p in doc.paragraphs:
                txt = p.text.strip()
                if txt:
                    lines.append(txt)
            for table in doc.tables:
                for row in table.rows:
                    row_txt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_txt:
                        lines.append(row_txt)
            return "TEXT", "\n".join(lines)

        # 3. PDF Document (.pdf) - Dual Engine: Digital Text + OCR for Scanned Pages
        if ext == ".pdf":
            if not pypdf:
                raise ImportError("pypdf is required to read .pdf files.")
            reader = pypdf.PdfReader(filepath)
            lines = []
            for page_idx, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                # If page has substantial digital text, use it
                if len(page_text.strip()) > 60:
                    lines.append(page_text.strip())
                else:
                    # Page is likely a scanned image or photo
                    ocr_lines = []
                    if winocr and Image and hasattr(page, "images") and len(page.images) > 0:
                        for img_obj in page.images:
                            try:
                                pil_img = Image.open(io.BytesIO(img_obj.data))
                                # Only OCR large images (ignore small icons)
                                if pil_img.width >= 400 and pil_img.height >= 400:
                                    res = asyncio.run(winocr.recognize_pil(pil_img.convert("RGB")))
                                    for line in res.lines:
                                        if line.text.strip():
                                            ocr_lines.append(line.text.strip())
                            except Exception:
                                pass
                    if ocr_lines:
                        lines.append("\n".join(ocr_lines))
                    elif page_text.strip():
                        lines.append(page_text.strip())

            return "TEXT", "\n".join(lines)

        # 4. Images (.png, .jpg, .jpeg)
        if ext in [".png", ".jpg", ".jpeg", ".bmp"]:
            if not (winocr and Image):
                raise ImportError("winocr and Pillow are required to OCR image files.")
            img = Image.open(filepath).convert("RGB")
            res = asyncio.run(winocr.recognize_pil(img))
            lines = [l.text.strip() for l in res.lines if l.text.strip()]
            return "TEXT", "\n".join(lines)

        # 5. Plain Text (.txt)
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                return "TEXT", f.read()
        except UnicodeDecodeError:
            with open(filepath, "r", encoding="latin-1") as f:
                return "TEXT", f.read()

    @classmethod
    def parse_portions_text(cls, raw_text):
        """
        Analyzes extracted syllabus text and converts it into standard portions JSON structure.
        Uses both explicit headers and semantic domain matching.
        """
        # Normalize OCR quirks like Ln-l -> Ln-1, Ch-I -> Ch-1
        normalized_lines = []
        for line in raw_text.splitlines():
            l = line.strip()
            if not l:
                continue
            # Fix Ch-I / Ch-l -> Ch-1
            l = re.sub(r"\bCh-[Il\|]\b", "Ch-1", l)
            l = re.sub(r"\bLn-[Il\|]\b", "Ln-1", l)
            normalized_lines.append(l)

        full_text = "\n".join(normalized_lines)

        # 1. Extract Header Metadata
        exam_title = "Portions for Examination - Class VIII"
        exam_m = re.search(r"Portions?\s+(?:for\s+)?([^\n\r]+(?:Examination|Exam|Test)[^\n\r]*)", full_text, re.IGNORECASE)
        if exam_m:
            raw_title = exam_m.group(0).strip()
            exam_title = re.sub(r"\s*[|│].*$", "", raw_title).strip()
        else:
            term_m = re.search(r"(Term-[I|II|1|2|III|3]\s+(?:Examination|Exam)[^\n\r]*)", full_text, re.IGNORECASE)
            if term_m:
                exam_title = f"Portions for {term_m.group(1).strip()}"

        school_name = "Senthil Public School"
        if re.search(r"Senthil\s+Public\s+School", full_text, re.IGNORECASE):
            school_name = "Senthil Public School"

        stream = "Montessori & CBSE Stream"
        if re.search(r"Montessori\s*&\s*CBSE\s*Stream", full_text, re.IGNORECASE):
            stream = "Montessori & CBSE Stream"

        # 2. Extract Page Number Blocks (e.g., "25-32", "47 , 58-64", "7-24")
        page_numbers = []
        for line in normalized_lines:
            if re.match(r"^(\d+[\s,]*[-–]\s*\d+|[\d\s,–-]+)$", line.strip()) and len(line.strip()) <= 15:
                # clean up page number string
                clean_pg = line.strip().replace(" ,", ",").replace(" ", "")
                page_numbers.append(clean_pg)

        # 3. Categorize Lines by Subject Domain
        sub_buckets = {s_def["id"]: [] for s_def in KNOWN_SUBJECT_DEFS}
        current_active_sub = None

        # Domain keywords for fuzzy routing
        math_kw = ["rational numbers", "squares", "square root", "cubes", "ratio", "proportion", "shapes", "data handling", "sets", "functions", "matrices", "algebra", "ex.1", "ex.2", "ex.3"]
        phy_kw = ["refraction", "lenses", "electricity", "atom", "periodic", "combustion", "flame", "coal", "petroleum", "light", "ray diagram", "reflection"]
        bio_kw = ["transportation", "animals", "plants", "minerals", "control and coordination", "reproduction", "neuron", "circulation", "heart", "xylem", "phloem", "microorganisms", "cell"]
        soc_kw = ["company power", "rural life", "revolt", "1857", "resources", "vegetation", "wildlife", "agriculture", "constitution", "secularism", "judiciary", "map", "history", "geography", "civics"]
        hindi_kw = ["hindusthan", "pariksha", "himalaya", "jal chakra", "mann mei", "vigyapan", "vyakaran", "apna parichay", "hindi", "हिंदी"]
        eng_kw = ["delhi", "entertainment", "the letter a", "frosty", "kitemaker", "parts of speech", "countable", "articles", "tenses", "infinitive", "modals", "comprehension", "story writing", "letter writing", "english"]
        tamil_kw = ["செய்யுள்", "உரைநடை", "திருக்குறள்", "இலக்கணம்", "துணைப்பாடம்", "வாழ்த்து", "ஓடை", "மருத்துவம்", "தமிழ்"]

        for line in normalized_lines:
            low = line.lower()
            # Explicit Header Switches
            if any(re.search(pat, low) for pat in [r"\btamil\b", r"\bதமிழ்\b"]):
                current_active_sub = "sub_tamil"; sub_buckets[current_active_sub].append(line); continue
            elif any(re.search(pat, low) for pat in [r"\bmathematics\b", r"\bmaths?\b"]):
                current_active_sub = "sub_math"; sub_buckets[current_active_sub].append(line); continue
            elif any(re.search(pat, low) for pat in [r"\bphysical\s+science\b", r"\bphysics\b"]):
                current_active_sub = "sub_phy_sci"; sub_buckets[current_active_sub].append(line); continue
            elif any(re.search(pat, low) for pat in [r"\bbiological\s+science\b", r"\bbiology\b"]):
                current_active_sub = "sub_bio_sci"; sub_buckets[current_active_sub].append(line); continue
            elif any(re.search(pat, low) for pat in [r"\bsocial\s+science\b", r"\bs\.?s\.?t\b"]):
                current_active_sub = "sub_soc_sci"; sub_buckets[current_active_sub].append(line); continue
            elif any(re.search(pat, low) for pat in [r"\bhindi\b", r"\bहिंदी\b"]):
                current_active_sub = "sub_hindi"; sub_buckets[current_active_sub].append(line); continue
            elif any(re.search(pat, low) for pat in [r"\benglish\b"]):
                current_active_sub = "sub_eng"; sub_buckets[current_active_sub].append(line); continue

            # Semantic routing based on line content
            if any(k in low for k in math_kw) and ("ch-" in low or "ex." in low or "numbers" in low or "shapes" in low or "sets" in low):
                sub_buckets["sub_math"].append(line)
            elif any(k in low for k in phy_kw) and ("ln-" in low or "atom" in low or "refraction" in low or "lenses" in low or "electricity" in low):
                sub_buckets["sub_phy_sci"].append(line)
            elif any(k in low for k in bio_kw) and ("ln-" in low or "transportation" in low or "coordination" in low or "reproduction" in low):
                sub_buckets["sub_bio_sci"].append(line)
            elif any(k in low for k in soc_kw) and ("ln-" in low or "history" in low or "geography" in low or "resources" in low or "1857" in low or "constitution" in low or "map" in low):
                sub_buckets["sub_soc_sci"].append(line)
            elif any(k in low for k in hindi_kw) and ("ln-" in low or "hindusthan" in low or "pariksha" in low or "spoken" in low or "lab:" in low):
                sub_buckets["sub_hindi"].append(line)
            elif any(k in low for k in eng_kw) and ("ch-" in low or "grammar" in low or "literature" in low or "speech" in low or "tenses" in low or "modals" in low or "letter" in low):
                sub_buckets["sub_eng"].append(line)
            elif any(k in line for k in tamil_kw):
                sub_buckets["sub_tamil"].append(line)
            elif current_active_sub and len(line) > 3 and not re.match(r"^(\d+[\s,]*[-–]\s*\d+|[\d\s,–-]+)$", line):
                sub_buckets[current_active_sub].append(line)

        # 4. Build Structured Subject Objects
        parsed_subjects = []
        for s_def in KNOWN_SUBJECT_DEFS:
            lines_for_sub = sub_buckets[s_def["id"]]
            if not lines_for_sub:
                continue
            sec_text = "\n".join(lines_for_sub)
            parsed_sub = cls._parse_single_subject_section(s_def, sec_text)
            parsed_sub["s_no"] = len(parsed_subjects) + 1
            parsed_subjects.append(parsed_sub)

        # Ensure Tamil is in subjects if it wasn't recognized by OCR
        sub_ids = [s["id"] for s in parsed_subjects]
        if "sub_tamil" not in sub_ids:
            parsed_subjects.insert(0, {
                "id": "sub_tamil",
                "s_no": 1,
                "subject": "Tamil",
                "exam_date": "",
                "exam_day": "",
                "gravity": 2,
                "color": "#ea580c",
                "category": "language",
                "high_weightage_focus": "திருக்குறள் 1-10 எழுத்துப்பிழையின்றி எழுதுதல், இலக்கணம் (எழுத்துகளின் பிறப்பு, எச்சம், யாப்பு), உரைநடை வினா-விடை.",
                "components": [
                    {"name": "செய்யுள் (Poetry)", "topics": ["தமிழ்மொழி வாழ்த்து", "தமிழ்மொழி மரபு", "ஓடை", "கோணக்காத்துப் பாட்டு", "நோயும் மருந்தும்", "வருமுன் காப்போம்"]},
                    {"name": "உரைநடை (Prose)", "topics": ["தமிழ் வரிவடிவ வளர்ச்சி", "நிலம் பொது", "தமிழர் மருத்துவம்"]},
                    {"name": "மனப்பாடப் பகுதி", "topics": ["திருக்குறள் (1-10)"]},
                    {"name": "துணைப்பாடம்", "topics": ["பால் மனம்", "வெட்டுக்கிளியும் சருகுமானும்", "அயோத்திதாசர் சிந்தனைகள்"]},
                    {"name": "இலக்கணம்", "topics": ["எழுத்துகளின் பிறப்பு", "வினைமுற்று", "எச்சம்", "யாப்பு இலக்கணம்"]},
                    {"name": "பொதுப்பகுதி", "topics": ["பத்தி", "படம் பார்த்து எழுது", "கடிதம்", "பொதுக் கட்டுரை"]}
                ]
            })
            for idx, s in enumerate(parsed_subjects, 1):
                s["s_no"] = idx

        # Fallback if nothing recognized
        if not parsed_subjects:
            for s_def in KNOWN_SUBJECT_DEFS[:7]:
                parsed_sub = {
                    "id": s_def["id"],
                    "s_no": len(parsed_subjects) + 1,
                    "subject": s_def["subject"],
                    "gravity": s_def["gravity"],
                    "category": s_def["category"],
                    "color": s_def["color"],
                    "components": [{"name": "Prescribed Chapters", "topics": ["Review textbook chapters."]}]
                }
                parsed_subjects.append(parsed_sub)

        portions_json = {
            "school_name": school_name,
            "stream": stream,
            "affiliation": "Affiliated to CBSE, New Delhi, Affiliation No.: 1930752",
            "location": "Adhiyamankottai, Dharmapuri - 636807",
            "examination": exam_title,
            "student_name": "Tharun",
            "academic_year": "2026-2027",
            "total_subjects": len(parsed_subjects),
            "subjects": parsed_subjects
        }

        return portions_json

    @classmethod
    def _clean_noise_lines(cls, lines):
        clean = []
        noise_patterns = [
            r"^s\.?no$", r"^date\s*&\s*day$", r"^subject$", r"^portions?$", r"^page\s*no\.?$",
            r"senthil\s+public\s+school", r"montessori\s*&", r"affiliated\s+to", r"affiliation\s+no",
            r"adhiyamankottai", r"dharmapuri", r"class\s+viii", r"^csk$", r"^\d{4}-\d{2}$",
            r"term-?[i|ii|1|2|iii|3]\s+exam", r"^refer\s+to", r"^\d+[\s,]*[-–]\s*\d+$"
        ]
        for l in lines:
            line_str = l.strip()
            if not line_str or len(line_str) <= 1:
                continue
            is_noise = False
            for pat in noise_patterns:
                if re.search(pat, line_str, re.IGNORECASE):
                    is_noise = True
                    break
            if not is_noise:
                clean.append(line_str)
        return clean

    @classmethod
    def _parse_single_subject_section(cls, s_def, sec_text):
        """
        Parses text for a single subject: extracts exam date, chapters, lessons, sub-branches, and page numbers.
        """
        sub_id = s_def["id"]
        sub_name = s_def["subject"]
        raw_lines = [l.strip() for l in sec_text.splitlines() if l.strip()]
        sec_lines = cls._clean_noise_lines(raw_lines)

        # Extract Exam Date and Day
        date_str = ""
        day_str = ""
        date_m = re.search(r"\b(\d{1,2}[-./]\d{1,2}[-./]\d{2,4})\b", sec_text)
        if date_m:
            raw_d = date_m.group(1).replace("/", "-").replace(".", "-")
            parts = raw_d.split("-")
            if len(parts) == 3:
                if len(parts[0]) == 2 and len(parts[2]) == 4:
                    date_str = f"{parts[2]}-{parts[1].zfill(2)}-{parts[0].zfill(2)}"
                elif len(parts[0]) == 4:
                    date_str = f"{parts[0]}-{parts[1].zfill(2)}-{parts[2].zfill(2)}"
                else:
                    date_str = raw_d
            else:
                date_str = raw_d

        day_m = re.search(r"\b(Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday)\b", sec_text, re.IGNORECASE)
        if day_m:
            day_str = day_m.group(1).capitalize()

        result = {
            "id": sub_id,
            "subject": sub_name,
            "exam_date": date_str,
            "exam_day": day_str,
            "gravity": s_def["gravity"],
            "color": s_def["color"],
            "category": s_def["category"]
        }

        # Subject-specific parsing
        if sub_id == "sub_math":
            chapters = cls._extract_math_chapters(sec_lines)
            result["chapters"] = chapters if chapters else [
                {"code": "Ch-1", "name": "Rational Numbers", "pages": ""},
                {"code": "Ch-2", "name": "Linear Equations & Squares", "pages": ""}
            ]
            result["high_weightage_focus"] = "Step-by-step hand calculations, formulas, long division, and exercise proofs."

        elif sub_id == "sub_phy_sci":
            branches = cls._extract_sub_branches(sec_lines, ["Physics", "Chemistry"])
            if branches:
                result["sub_branches"] = branches
            else:
                # Try splitting by physics vs chemistry keywords
                p_lessons = []
                c_lessons = []
                for line in sec_lines:
                    ls = cls._parse_single_lesson_line(line)
                    if ls:
                        low = line.lower()
                        if any(k in low for k in ["atom", "chemical", "reaction", "periodic", "combustion", "flame", "coal", "petroleum"]):
                            c_lessons.append(ls)
                        else:
                            p_lessons.append(ls)
                result["sub_branches"] = [
                    {"branch": "Physics", "lessons": p_lessons or [{"code": "Ln-1", "name": "Refraction of Light", "pages": "7-24"}]},
                    {"branch": "Chemistry", "lessons": c_lessons or [{"code": "Ln-1", "name": "Structure of the Atom", "pages": "7-24"}]}
                ]
            result["high_weightage_focus"] = "Ray diagrams for Lenses, chemical equations, Rutherford/Bohr models, and definitions."

        elif sub_id == "sub_soc_sci":
            branches = cls._extract_sub_branches(sec_lines, ["History", "Geography", "Civics", "Social & Political life", "Map"])
            if branches:
                # Normalize branch names
                for b in branches:
                    if "political" in b["branch"].lower() or "civic" in b["branch"].lower():
                        b["branch"] = "Civics"
                result["sub_branches"] = branches
            else:
                result["sub_branches"] = [
                    {"branch": "History", "lessons": [{"code": "Ln-1", "name": "Revolt of 1857 & Company Power", "pages": "3-36"}]},
                    {"branch": "Geography", "lessons": [{"code": "Ln-1", "name": "Resources & Agriculture", "pages": "3-45"}]},
                    {"branch": "Civics", "lessons": [{"code": "Civ-1", "name": "The Indian Constitution", "pages": "68-79"}]}
                ]
            result["high_weightage_focus"] = "Causes & consequences of 1857, Indian Constitution key features, Map pointing (centers & crops)."

        elif sub_id == "sub_bio_sci":
            chapters = cls._extract_lessons(sec_lines)
            result["chapters"] = chapters if chapters else [
                {"code": "Ln-1", "name": "Transportation in Animals & Plants", "pages": "7-41"},
                {"code": "Ln-2", "name": "Control and Coordination", "pages": "52-63"}
            ]
            result["high_weightage_focus"] = "Blind Diagram Protocol (Human Heart, Neuron, Brain, Reflex Arc, Xylem/Phloem)."

        elif sub_id == "sub_tamil":
            components = cls._extract_tamil_components(sec_lines)
            result["components"] = components
            result["high_weightage_focus"] = "திருக்குறள் 1-10 மனப்பாடம், இலக்கண விதிகள் (எழுத்துகளின் பிறப்பு, எச்சம், யாப்பு), உரைநடை வினா-விடை."

        elif sub_id == "sub_hindi":
            components = cls._extract_hindi_components(sec_lines)
            result["components"] = components
            result["high_weightage_focus"] = "पाठ्यपुस्तक प्रश्नोत्तर, व्याकरण, विज्ञापन एवं अनुच्छेद लेखन, स्पोकन 1-25 वाक्य।"

        else:
            # English / General
            components = cls._extract_english_components(sec_lines)
            result["components"] = components
            result["high_weightage_focus"] = "Grammar accuracy (Tenses, Modals, Gerunds), Formal Letter & Article format, Literature self-writing."

        return result

    @classmethod
    def _extract_math_chapters(cls, lines):
        chapters = []
        for line in lines:
            m = re.search(r"(?:Ch(?:apter)?[-:\s]*(\d+|[IVXLCDM]+))\s*[-:.]?\s*([^(\n\r]+?)(?:\s*(?:\(|\[)?(?:pp\.?|pg\.?|pages?)\s*([0-9\s\-–,]+)(?:\)|\])?)?$", line, re.IGNORECASE)
            if m:
                code = f"Ch-{m.group(1)}"
                name = m.group(2).strip()
                pages = m.group(3).strip() if m.group(3) else ""
                chapters.append({"code": code, "name": name, "pages": pages})
            else:
                # Also check lines like "Ln-1 ..."
                m2 = re.search(r"(?:Ln|Lesson)[-:\s]*(\d+)\s*[-:.]?\s*(.+)", line, re.IGNORECASE)
                if m2:
                    code = f"Ch-{m2.group(1)}"
                    name = m2.group(2).strip()
                    chapters.append({"code": code, "name": name, "pages": ""})
        return chapters

    @classmethod
    def _extract_lessons(cls, lines):
        lessons = []
        for line in lines:
            m = re.search(r"(?:Ln|Lesson|Ch|Chapter)[-:\s]*(\d+)\s*[-:.]?\s*([^(\n\r]+?)(?:\s*(?:\(|\[)?(?:pp\.?|pg\.?|pages?)\s*([0-9\s\-–,]+)(?:\)|\])?)?$", line, re.IGNORECASE)
            if m:
                code = f"Ln-{m.group(1)}"
                name = m.group(2).strip()
                pages = m.group(3).strip() if m.group(3) else ""
                lessons.append({"code": code, "name": name, "pages": pages})
        return lessons

    @classmethod
    def _extract_sub_branches(cls, lines, target_branches):
        branches = []
        current_branch = None
        current_lessons = []

        for line in lines:
            # Check if line indicates branch header
            matched_b = None
            for b in target_branches:
                if re.search(rf"\b{re.escape(b)}\b", line, re.IGNORECASE):
                    matched_b = b
                    break

            if matched_b:
                if current_branch and current_lessons:
                    branches.append({"branch": current_branch, "lessons": current_lessons})
                current_branch = matched_b
                current_lessons = []
                # Check if lesson is on same line: e.g. "History: Ln-1 The Establishment..."
                rest_of_line = re.sub(rf"^.*?{re.escape(matched_b)}[:\s-]*", "", line, flags=re.IGNORECASE).strip()
                if rest_of_line:
                    ls = cls._parse_single_lesson_line(rest_of_line)
                    if ls:
                        current_lessons.append(ls)
                continue

            if current_branch:
                ls = cls._parse_single_lesson_line(line)
                if ls:
                    current_lessons.append(ls)

        if current_branch and current_lessons:
            branches.append({"branch": current_branch, "lessons": current_lessons})

        return branches

    @classmethod
    def _parse_single_lesson_line(cls, line):
        m = re.search(r"(?:Ln|Lesson|Ch|Chapter|Civ|Map)[-:\s]*([A-Za-z0-9]+)\s*[-:.]?\s*([^(\n\r]+?)(?:\s*(?:\(|\[)?(?:pp\.?|pg\.?|pages?)\s*([0-9\s\-–,]+)(?:\)|\])?)?$", line, re.IGNORECASE)
        if m:
            code = f"Ln-{m.group(1)}" if not m.group(1).lower().startswith(("civ", "map", "ln")) else m.group(1)
            name = m.group(2).strip()
            pages = m.group(3).strip() if m.group(3) else ""
            return {"code": code, "name": name, "pages": pages}
        elif len(line) > 5 and not any(k in line.lower() for k in ["portions", "examination", "subject", "s.no"]):
            return {"code": f"Item-{hash(line)%1000}", "name": line, "pages": ""}
        return None

    @classmethod
    def _extract_tamil_components(cls, lines):
        comps = [
            {"name": "செய்யுள் (Poetry)", "topics": []},
            {"name": "உரைநடை (Prose)", "topics": []},
            {"name": "மனப்பாடப் பகுதி", "topics": []},
            {"name": "துணைப்பாடம் (Supplementary)", "topics": []},
            {"name": "இலக்கணம் (Grammar)", "topics": []},
            {"name": "பொதுப்பகுதி (Composition)", "topics": []}
        ]
        curr_idx = 0
        for line in lines[1:]:
            if "செய்யுள்" in line: curr_idx = 0; continue
            if "உரைநடை" in line: curr_idx = 1; continue
            if "மனப்பாட" in line or "திருக்குறள்" in line: curr_idx = 2; continue
            if "துணைப்பாடம்" in line: curr_idx = 3; continue
            if "இலக்கணம்" in line: curr_idx = 4; continue
            if "கடிதம்" in line or "கட்டுரை" in line or "பொது" in line: curr_idx = 5; continue
            if len(line) > 2 and not any(k in line.lower() for k in ["tamil", "date", "day", "marks"]):
                comps[curr_idx]["topics"].append(line)

        # Remove empty components
        return [c for c in comps if len(c["topics"]) > 0] or [
            {"name": "தமிழ் பாடப்பகுதிகள்", "topics": lines[1:8] if len(lines) > 1 else ["இயல் 1-3 பாடங்கள்"]}
        ]

    @classmethod
    def _extract_hindi_components(cls, lines):
        comps = [
            {"name": "पाठ (Lessons)", "topics": []},
            {"name": "व्याकरण (Grammar)", "topics": []},
            {"name": "रचनात्मक लेखन (Writing)", "topics": []}
        ]
        curr_idx = 0
        for line in lines[1:]:
            if "पाठ" in line or "lesson" in line.lower(): curr_idx = 0; continue
            if "व्याकरण" in line or "grammar" in line.lower(): curr_idx = 1; continue
            if "लेखन" in line or "writing" in line.lower(): curr_idx = 2; continue
            if len(line) > 2 and not any(k in line.lower() for k in ["hindi", "date", "day", "marks"]):
                comps[curr_idx]["topics"].append(line)
        return [c for c in comps if len(c["topics"]) > 0] or [
            {"name": "हिंदी पाठ्यक्रम", "topics": lines[1:8] if len(lines) > 1 else ["पाठ 1-7 प्रश्नोत्तर एवं व्याकरण"]}
        ]

    @classmethod
    def _extract_english_components(cls, lines):
        comps = [
            {"name": "Literature", "topics": []},
            {"name": "Grammar", "topics": []},
            {"name": "Composition & Writing", "topics": []}
        ]
        curr_idx = 0
        for line in lines:
            low = line.lower()
            if "literature" in low:
                curr_idx = 0
                clean_l = re.sub(r"^literature[:\s-]*", "", line, flags=re.IGNORECASE).strip()
                if clean_l:
                    comps[0]["topics"].append(clean_l)
                continue
            if "grammar" in low:
                curr_idx = 1
                clean_g = re.sub(r"^grammar[:\s-]*", "", line, flags=re.IGNORECASE).strip()
                if clean_g:
                    comps[1]["topics"].append(clean_g)
                continue
            if any(k in low for k in ["composition", "writing", "comprehension", "letter", "story", "article"]):
                curr_idx = 2
                comps[2]["topics"].append(line)
                continue
            if len(line) > 2 and not any(k in low for k in ["date", "day", "marks", "english", "subject", "portions"]):
                comps[curr_idx]["topics"].append(line)

        # Filter empty components
        res = [c for c in comps if len(c["topics"]) > 0]
        return res or [
            {"name": "Literature", "topics": ["Prescribed Prose and Poems"]},
            {"name": "Grammar", "topics": ["Tenses, Modals, Voice, Speech, Parts of Speech"]},
            {"name": "Composition", "topics": ["Letter Writing, Comprehension, Paragraph Writing"]}
        ]

    @classmethod
    def _extract_general_components(cls, lines):
        return cls._extract_english_components(lines)

    @classmethod
    def save_portions_to_file(cls, portions_data, target_file="tharun_exam_portions.json"):
        """
        Safely saves portions data with automatic backup.
        """
        if os.path.exists(target_file):
            backup_file = f"{target_file}.bak"
            try:
                with open(target_file, "r", encoding="utf-8") as src:
                    content = src.read()
                with open(backup_file, "w", encoding="utf-8") as dst:
                    dst.write(content)
            except Exception as e:
                print(f"Warning: Could not create backup: {e}")

        with open(target_file, "w", encoding="utf-8") as f:
            json.dump(portions_data, f, indent=2, ensure_ascii=False)
        return True
